"""Importación de extractos bancarios (CSV o Excel).

Dos pasos sin guardar nada en el medio:
1. `preview`: lee el archivo, detecta las columnas, marca posibles duplicados y sugiere categorías.
2. `commit`: el usuario confirma (y corrige) las filas y se crean los movimientos.
"""

import csv
import io
import re
import zipfile
from collections import Counter
from datetime import date, datetime, timedelta
from decimal import Decimal

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.errors import PayloadTooLargeError, ValidationError
from app.models import CategoryKind, Transaction, TransactionKind, User
from app.schemas.common import MAX_YEAR, MIN_YEAR, clean_text
from app.schemas.support import ImportCommit, ImportMapping, ImportPreview, ImportResult, ImportRow
from app.schemas.transaction import TransactionIn
from app.services import accounts, rules, transactions
from app.services.invoice_parser import MAX_AMOUNT, parse_amount

MAX_IMPORT_BYTES = 5 * 1024 * 1024
MAX_UNZIPPED_BYTES = 60 * 1024 * 1024  # un .xlsx es un zip: se limita para evitar bombas de descompresión
MAX_ROWS = 5000
# Un extracto tiene pocas columnas. Sin tope, un Excel que declara 16.384 columnas por fila
# obligaba a armar millones de celdas vacías en memoria.
MAX_COLS = 60
PREVIEW_ROWS = 2000
HEADER_SCAN_ROWS = 25

_HEADERS = {
    "date": ("fecha", "date", "f. valor", "fec."),
    "description": ("descrip", "concepto", "detalle", "movimiento", "referencia", "glosa", "comercio"),
    "debit": ("débito", "debito", "debe", "cargo", "egreso", "retiro", "salida"),
    "credit": ("crédito", "credito", "haber", "abono", "ingreso", "depósito", "deposito", "entrada"),
    "amount": ("importe", "monto", "valor", "amount"),
}
_DATE_PATTERNS = (
    (re.compile(r"^(\d{4})-(\d{1,2})-(\d{1,2})"), "ymd"),
    (re.compile(r"^(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{4})"), "dmy"),
    (re.compile(r"^(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{2})$"), "dmy2"),
)


# ---------- Lectura del archivo ----------


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        return format(Decimal(str(value)), "f")
    return clean_text(str(value)).strip()


def _read_xlsx(content: bytes) -> list[list[str]]:
    from openpyxl import load_workbook  # import local: solo hace falta al importar Excel

    with zipfile.ZipFile(io.BytesIO(content)) as zf:
        if sum(i.file_size for i in zf.infolist()) > MAX_UNZIPPED_BYTES:
            raise ValidationError("El Excel es demasiado grande una vez descomprimido")
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        return [[_cell(c) for c in row] for row in ws.iter_rows(max_row=MAX_ROWS, max_col=MAX_COLS, values_only=True)]
    finally:
        wb.close()


def _read_csv(content: bytes) -> list[list[str]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("latin-1")
    # El separador es el que más se repite en las primeras líneas. (csv.Sniffer falla cuando el
    # extracto empieza con un título como "Banco X - Extracto de cuenta".)
    sample = text[:16384].splitlines()[:60]
    delimiter = max(";,\t|", key=lambda d: sum(line.count(d) for line in sample))
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows = []
    for i, row in enumerate(reader):
        if i >= MAX_ROWS:
            break
        rows.append([clean_text(c).strip() for c in row[:MAX_COLS]])
    return rows


def read_table(content: bytes) -> list[list[str]]:
    if len(content) > MAX_IMPORT_BYTES:
        raise PayloadTooLargeError("El archivo supera 5 MB")
    if content.startswith(b"PK\x03\x04"):
        return _read_xlsx(content)
    if content.startswith(b"\xd0\xcf\x11\xe0"):
        raise ValidationError("Es un Excel antiguo (.xls). Guardalo como .xlsx o .csv y volvé a probar.")
    if b"\x00" in content[:1024]:
        raise ValidationError("Formato no soportado. Usá un archivo .csv o .xlsx")
    return _read_csv(content)


# ---------- Columnas ----------


def detect_mapping(rows: list[list[str]]) -> ImportMapping | None:
    best: tuple[int, int, dict[str, int]] | None = None
    for r, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        found: dict[str, int] = {}
        for c, cell in enumerate(row):
            low = cell.lower()
            for field, words in _HEADERS.items():
                if field not in found and any(w in low for w in words):
                    found[field] = c
                    break
        score = len(found)
        if "date" in found and score >= 2 and (best is None or score > best[1]):
            best = (r, score, found)
    if best is None:
        return None
    header_row, _, found = best
    has_split = "debit" in found or "credit" in found
    return ImportMapping(
        header_row=header_row,
        date=found["date"],
        description=found.get("description"),
        amount=None if has_split else found.get("amount"),
        debit=found.get("debit"),
        credit=found.get("credit"),
    )


def _parse_date(value: str, day_first: bool) -> date | None:
    found = _read_date(value.strip(), day_first)
    # Fuera del rango que acepta la app: se marca como fecha inválida en la vista previa.
    return found if found and MIN_YEAR <= found.year <= MAX_YEAR else None


def _read_date(value: str, day_first: bool) -> date | None:
    for pattern, kind in _DATE_PATTERNS:
        if not (m := pattern.match(value)):
            continue
        a, b, c = (int(x) for x in m.groups())
        try:
            if kind == "ymd":
                return date(a, b, c)
            year = c if kind == "dmy" else 2000 + c
            return date(year, b, a) if day_first else date(year, a, b)
        except ValueError:
            return None
    if re.fullmatch(r"[0-9]{5}(\.0+)?", value):  # número de serie de Excel
        return date(1899, 12, 30) + timedelta(days=int(float(value)))
    return None


def _signed_amount(value: str) -> Decimal | None:
    value = value.strip()
    if not value:
        return None
    negative = value.startswith("-") or value.endswith("-") or (value.startswith("(") and value.endswith(")"))
    amount = parse_amount(value)
    if amount is None or amount > MAX_AMOUNT:
        return None
    return -amount if negative else amount


def _get(row: list[str], index: int | None) -> str:
    return row[index] if index is not None and index < len(row) else ""


# ---------- Vista previa y confirmación ----------


def preview(db: Session, user: User, content: bytes, mapping: ImportMapping | None) -> ImportPreview:
    rows = read_table(content)
    if not rows:
        raise ValidationError("El archivo está vacío")
    mapping = mapping or detect_mapping(rows)
    width = max(len(r) for r in rows[:HEADER_SCAN_ROWS])
    header = rows[mapping.header_row] if mapping else []
    columns = [header[i] if i < len(header) and header[i] else f"Columna {i + 1}" for i in range(width)]
    if mapping is None:
        # Sin encabezados reconocibles: el usuario elige las columnas en la pantalla.
        return ImportPreview(columns=columns, mapping=None, rows=[], skipped=0)
    if mapping.amount is None and mapping.debit is None and mapping.credit is None:
        raise ValidationError("Elegí la columna del monto (o las de débito y crédito)")

    suggester = rules.Suggester(db, user)
    parsed: list[ImportRow] = []
    skipped = 0
    for line, row in enumerate(rows[mapping.header_row + 1 :], start=mapping.header_row + 2):
        if not any(cell.strip() for cell in row):
            continue
        raw_date = _get(row, mapping.date)
        when = _parse_date(raw_date, mapping.day_first)
        description = _get(row, mapping.description)[:160]
        if mapping.amount is not None:
            signed = _signed_amount(_get(row, mapping.amount))
        else:
            debit = _signed_amount(_get(row, mapping.debit))
            credit = _signed_amount(_get(row, mapping.credit))
            signed = -abs(debit) if debit else (abs(credit) if credit else None)
        if when is None and signed is None:
            skipped += 1  # filas de totales, saldos o texto suelto
            continue
        kind = TransactionKind.expense if signed is not None and signed < 0 else TransactionKind.income
        error = None
        if when is None:
            error = f"Fecha no reconocida: «{raw_date}»"
        elif not signed:
            error = "Sin monto"
        suggestion = suggester.suggest(CategoryKind(kind.value), text=description)
        parsed.append(
            ImportRow(
                line=line,
                date=when,
                description=description,
                amount=abs(signed) if signed else None,
                kind=kind,
                category_id=suggestion.category_id,
                duplicate=False,
                error=error,
            )
        )
        if len(parsed) >= PREVIEW_ROWS:
            break
    _mark_duplicates(db, user, parsed)
    return ImportPreview(columns=columns, mapping=mapping, rows=parsed, skipped=skipped)


def _mark_duplicates(db: Session, user: User, rows: list[ImportRow]) -> None:
    dated = [r.date for r in rows if r.date]
    if not dated:
        return
    existing = db.execute(
        select(Transaction.date, Transaction.amount, Transaction.kind).where(
            Transaction.user_id == user.id, Transaction.date.between(min(dated), max(dated))
        )
    )
    # Se cuenta cuántos hay de cada (fecha, monto, tipo): si el extracto trae dos compras iguales el
    # mismo día y solo una estaba cargada, se marca una sola como posible repetida.
    seen = Counter((d, a, k) for d, a, k in existing)
    for r in rows:
        key = (r.date, r.amount, r.kind)
        if r.date and r.amount is not None and seen[key] > 0:
            r.duplicate = True
            seen[key] -= 1


def commit(db: Session, user: User, data: ImportCommit) -> ImportResult:
    account_id = accounts.resolve(db, user, data.account_id)
    total = Decimal(0)
    learned: set[str] = set()
    for item in data.items:
        tx = transactions.create(
            db,
            user,
            TransactionIn(
                kind=item.kind,
                amount=item.amount,
                date=item.date,
                description=item.description,
                method=item.method,
                category_id=item.category_id,
                account_id=account_id,
            ),
            learn=False,
        )
        total += tx.amount
        key = rules.normalize(item.description)
        if item.category_id and key not in learned:
            learned.add(key)
            rules.learn_from_text(db, user, item.description, tx.category_id)
    return ImportResult(created=len(data.items), total=total)
